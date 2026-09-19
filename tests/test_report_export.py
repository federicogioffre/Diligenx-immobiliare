"""Tests for GET /deals/{deal_id}/report/export.

These tests skip LLM generation entirely: they insert a `Report` row
directly and verify the HTTP behaviour of the export endpoint plus the
output produced by `render_pdf` / `render_docx`.
"""
from __future__ import annotations

import io

import pytest
from sqlalchemy.orm import Session

from app.core.security import create_access_token, hash_password
from app.models.deal import Deal
from app.models.report import Report, ReportStatus, ReportType
from app.models.user import User
from app.models.workspace import Workspace


SAMPLE_MARKDOWN = """\
## Executive Summary
Questa è una sintesi di test.

## Red Flags
- [CRITICAL] Clausola di cambio controllo
- [HIGH] Restrizioni al trasferimento delle quote
- [MEDIUM] Veto di minoranza

## Observations
Osservazione di prova con **testo in grassetto**.

## Information Requests
1. Documento uno
2. Documento due

## Documents Reviewed
| Document | Language | Pages |
|---|---|---|
| NDA.pdf | it | 3 |
| SPA.pdf | en | 42 |
"""


# ── Fixtures ───────────────────────────────────────────────────────────────


@pytest.fixture
def report_fixture(db: Session, test_deal: Deal) -> Report:
    """Insert a completed due-diligence report for ``test_deal``."""
    from tests.report_fixtures import complete_report_inputs
    fingerprint = complete_report_inputs(db, test_deal)
    r = Report(
        deal_id=test_deal.id,
        report_type=ReportType.due_diligence,
        title="Due Diligence Report",
        content=SAMPLE_MARKDOWN,
        status=ReportStatus.completed,
        input_hash=fingerprint,
    )
    db.add(r)
    db.commit()
    db.refresh(r)
    return r


def _make_other_user_with_report(db: Session) -> tuple[Workspace, Deal, Report, dict]:
    other = User(
        email="other_export@example.com",
        hashed_password=hash_password("pw"),
        full_name="Other",
    )
    db.add(other)
    db.commit()
    db.refresh(other)

    ws = Workspace(name="Other WS", owner_id=other.id)
    db.add(ws)
    db.commit()
    db.refresh(ws)

    deal = Deal(workspace_id=ws.id, title="Other Deal")
    db.add(deal)
    db.commit()
    db.refresh(deal)

    report = Report(
        deal_id=deal.id,
        report_type=ReportType.due_diligence,
        title="Other Report",
        content=SAMPLE_MARKDOWN,
        status=ReportStatus.completed,
    )
    db.add(report)
    db.commit()
    db.refresh(report)

    headers = {"Authorization": f"Bearer {create_access_token(other.email, user_id=other.id, hashed_password=other.hashed_password)}"}
    return ws, deal, report, headers


# ── HTTP behaviour ─────────────────────────────────────────────────────────


def test_export_requires_auth(client, test_deal):
    resp = client.get(f"/deals/{test_deal.id}/report/export?format=pdf")
    assert resp.status_code in (401, 403)


def test_export_invalid_format_returns_400(client, auth_headers, test_deal, report_fixture):
    resp = client.get(
        f"/deals/{test_deal.id}/report/export?format=rtf",
        headers=auth_headers,
    )
    assert resp.status_code == 400
    assert "pdf" in resp.json()["detail"].lower()


def test_export_404_when_no_report(client, auth_headers, test_deal):
    resp = client.get(
        f"/deals/{test_deal.id}/report/export?format=pdf",
        headers=auth_headers,
    )
    assert resp.status_code == 404


def test_export_404_for_deal_of_other_user(client, auth_headers, db):
    _, other_deal, _, _ = _make_other_user_with_report(db)
    resp = client.get(
        f"/deals/{other_deal.id}/report/export?format=pdf",
        headers=auth_headers,
    )
    assert resp.status_code == 404


def test_export_404_for_nonexistent_deal(client, auth_headers):
    resp = client.get(
        "/deals/9999/report/export?format=pdf",
        headers=auth_headers,
    )
    assert resp.status_code == 404


# ── PDF output ─────────────────────────────────────────────────────────────


def test_export_pdf_returns_pdf_bytes(client, auth_headers, test_deal, report_fixture):
    resp = client.get(
        f"/deals/{test_deal.id}/report/export?format=pdf",
        headers=auth_headers,
    )
    assert resp.status_code == 200
    assert resp.headers["content-type"].startswith("application/pdf")
    disposition = resp.headers.get("content-disposition", "")
    assert "attachment" in disposition
    assert f"deal-{test_deal.id}.pdf" in disposition
    # PDF magic header
    assert resp.content.startswith(b"%PDF-")
    # Sanity: should be at least a couple of KB
    assert len(resp.content) > 1024


# ── DOCX output ────────────────────────────────────────────────────────────


def test_export_docx_returns_docx_bytes(client, auth_headers, test_deal, report_fixture):
    resp = client.get(
        f"/deals/{test_deal.id}/report/export?format=docx",
        headers=auth_headers,
    )
    assert resp.status_code == 200
    assert resp.headers["content-type"].startswith(
        "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
    )
    disposition = resp.headers.get("content-disposition", "")
    assert f"deal-{test_deal.id}.docx" in disposition
    # DOCX is a ZIP → magic bytes "PK"
    assert resp.content[:2] == b"PK"
    assert len(resp.content) > 2048


def test_export_docx_contains_deal_title_and_sections(
    client, auth_headers, test_deal, report_fixture
):
    """Round-trip: re-open the generated DOCX with python-docx and inspect."""
    from docx import Document as _DocxDocument

    resp = client.get(
        f"/deals/{test_deal.id}/report/export?format=docx",
        headers=auth_headers,
    )
    assert resp.status_code == 200

    doc = _DocxDocument(io.BytesIO(resp.content))
    text = "\n".join(p.text for p in doc.paragraphs)

    # Branding + deal title in header block
    assert "DILIGENX" in text
    assert test_deal.title in text

    # Section headings are rendered as docx headings; they still appear in
    # paragraph text when iterating paragraphs (docx doesn't distinguish
    # headings from paragraphs at the iteration level).
    assert "Executive Summary" in text
    assert "Red Flags" in text
    assert "Documents Reviewed" in text

    # The Documents Reviewed pipe-table becomes a real Word table
    assert len(doc.tables) >= 1


# ── Landscape layout for the Key Findings section ──────────────────────────

FINDINGS_MARKDOWN = """\
## 1. Executive Summary
Sintesi di prova.

## 2. Key Findings

### 2.1 Corporate Matters

| # | Issue | Risk | Recommendation | Severity |
|---|-------|------|----------------|----------|
| 1 | Clausola di cambio di controllo nel contratto quadro | Risoluzione anticipata | Rinegoziare la clausola | [HIGH] |
| 2 | Quorum rafforzati in assemblea | Stallo decisionale | Patto parasociale | [MEDIUM] |

## 3. Observations
Osservazione successiva alla tabella.
"""


@pytest.fixture
def findings_report_fixture(db: Session, test_deal: Deal) -> Report:
    from tests.report_fixtures import complete_report_inputs
    fingerprint = complete_report_inputs(db, test_deal)
    r = Report(
        deal_id=test_deal.id,
        report_type=ReportType.due_diligence,
        title="Findings Report",
        content=FINDINGS_MARKDOWN,
        input_hash=fingerprint,
        status=ReportStatus.completed,
    )
    db.add(r)
    db.commit()
    db.refresh(r)
    return r


def test_pdf_key_findings_pages_are_landscape(
    client, auth_headers, test_deal, findings_report_fixture
):
    """The Key Findings section renders on landscape pages; the sections
    before and after stay portrait."""
    from pypdf import PdfReader

    resp = client.get(
        f"/deals/{test_deal.id}/report/export?format=pdf",
        headers=auth_headers,
    )
    assert resp.status_code == 200

    reader = PdfReader(io.BytesIO(resp.content))
    orientations = []
    findings_page_landscape = None
    for page in reader.pages:
        w, h = float(page.mediabox.width), float(page.mediabox.height)
        landscape = w > h
        orientations.append(landscape)
        if "Key Findings" in (page.extract_text() or ""):
            findings_page_landscape = landscape

    assert any(orientations), "expected at least one landscape page"
    assert not all(orientations), "cover/body pages must stay portrait"
    assert findings_page_landscape is True, (
        "the page carrying the Key Findings table must be landscape"
    )


def test_pdf_without_key_findings_stays_all_portrait(
    client, auth_headers, test_deal, report_fixture
):
    from pypdf import PdfReader

    resp = client.get(
        f"/deals/{test_deal.id}/report/export?format=pdf",
        headers=auth_headers,
    )
    assert resp.status_code == 200
    reader = PdfReader(io.BytesIO(resp.content))
    for page in reader.pages:
        w, h = float(page.mediabox.width), float(page.mediabox.height)
        assert h > w, "no landscape pages expected without a Key Findings section"


def test_docx_key_findings_section_is_landscape(
    client, auth_headers, test_deal, findings_report_fixture
):
    from docx import Document as _DocxDocument
    from docx.enum.section import WD_ORIENT

    resp = client.get(
        f"/deals/{test_deal.id}/report/export?format=docx",
        headers=auth_headers,
    )
    assert resp.status_code == 200

    doc = _DocxDocument(io.BytesIO(resp.content))
    sections = doc.sections
    # cover+summary (portrait) / key findings (landscape) / rest (portrait)
    assert len(sections) == 3
    assert sections[0].orientation == WD_ORIENT.PORTRAIT
    assert sections[1].orientation == WD_ORIENT.LANDSCAPE
    assert int(sections[1].page_width) > int(sections[1].page_height)
    assert sections[2].orientation == WD_ORIENT.PORTRAIT
    assert int(sections[2].page_width) < int(sections[2].page_height)


def test_docx_without_key_findings_single_portrait_section(
    client, auth_headers, test_deal, report_fixture
):
    from docx import Document as _DocxDocument

    resp = client.get(
        f"/deals/{test_deal.id}/report/export?format=docx",
        headers=auth_headers,
    )
    assert resp.status_code == 200
    doc = _DocxDocument(io.BytesIO(resp.content))
    assert len(doc.sections) == 1
