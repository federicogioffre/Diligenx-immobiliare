"""Corpus completeness, full-document batching and atomic ingestion regressions."""
import json
from types import SimpleNamespace
from unittest.mock import patch

import pytest
from docx import Document as DocxDocument

from app.models.agent_finding import AgentFinding
from app.models.document import Document, DocumentStatus
from app.models.document_chunk import DocumentChunk, EMBEDDING_DIM
from app.models.report import Report, ReportStatus, ReportType
from app.models.red_flag import RedFlag
from packages.agents.deal_analysis_agent import FakeLLMProvider, LLMResponseError
from packages.agents.document_coverage import aggregation_batches, collect_document_payloads, text_batches
from packages.agents.orchestrator import run_deal_analysis_pipeline
from packages.agents.red_flag_aggregator import RedFlagAggregationAgent
from packages.agents.report_agent import generate_due_diligence_report
from packages.real_estate.property_identity import consolidate_property_profiles
from packages.real_estate.property_identity_agent import PropertyIdentityAgent
from packages.real_estate.specialists import SPECIALISTS, EncumbrancesAgent
from packages.parsing.docx_parser import extract_docx
from packages.parsing.ingestion_pipeline import process_document
from packages.parsing.pdf_parser import PageText, extract_pdf


def document(db, deal, *, text="Source text", status=DocumentStatus.processed, path="/tmp/tax.txt"):
    doc = Document(deal_id=deal.id, filename="avviso_di_accertamento.txt", file_path=path,
                   mime_type="text/plain", file_size=100, status=status, raw_text=text)
    db.add(doc)
    db.commit()
    return doc


def test_docx_tables_nested_and_merged_cells_preserve_order(tmp_path):
    doc = DocxDocument()
    doc.add_paragraph("BEFORE")
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "DEBT"
    table.cell(0, 1).text = "EUR 250000"
    nested = table.cell(1, 0).add_table(rows=1, cols=1)
    nested.cell(0, 0).text = "NESTED GUARANTEE"
    table.cell(1, 1).text = "END TABLE"
    doc.add_paragraph("AFTER")
    merged = doc.add_table(rows=1, cols=2)
    merged.cell(0, 0).merge(merged.cell(0, 1)).text = "MERGED ONCE"
    path = tmp_path / "tables.docx"
    doc.save(path)

    text = extract_docx(path)[0].text
    expected = ["BEFORE", "DEBT", "EUR 250000", "NESTED GUARANTEE", "END TABLE", "AFTER"]
    assert [text.index(item) for item in expected] == sorted(text.index(item) for item in expected)
    assert text.count("MERGED ONCE") == 1


def test_unreadable_page_fails_a_partially_readable_pdf():
    reader = SimpleNamespace(pages=[SimpleNamespace(extract_text=lambda: "Readable first page"),
                                    SimpleNamespace(extract_text=lambda: "")])
    with patch("packages.parsing.pdf_parser.PdfReader", return_value=reader), \
         patch("packages.parsing.pdf_parser._ocr_page", return_value=""):
        with pytest.raises(ValueError, match="PDF page 2.*incomplete"):
            extract_pdf("unreadable.pdf")


@pytest.mark.parametrize("pages", [[], [PageText(1, "  ")]])
def test_empty_extraction_fails_and_preserves_previous_text(db, test_deal, pages):
    doc = document(db, test_deal, text="PREVIOUS VALID TEXT")
    with patch.dict("packages.parsing.ingestion_pipeline._EXTRACTORS", {"text/plain": lambda _: pages}), \
         patch("app.services.storage.open_local"):
        with pytest.raises(ValueError, match="no readable text"):
            process_document(doc.id, db)
    db.refresh(doc)
    assert doc.status == DocumentStatus.failed
    assert doc.raw_text == "PREVIOUS VALID TEXT"


def test_ingestion_has_no_transaction_during_extraction_and_embeddings(db, test_deal, tmp_path):
    path = tmp_path / "input.txt"
    path.write_text("Actual text", encoding="utf-8")
    doc = document(db, test_deal, path=str(path))

    def extract(_):
        assert not db.in_transaction()
        return [PageText(1, "Updated source text with a guarantee.")]

    def embed(texts):
        assert not db.in_transaction()
        assert texts == ["Updated source text with a guarantee."]
        return [[0.1] * EMBEDDING_DIM]

    provider = SimpleNamespace(embed_texts=embed)
    with patch.dict("packages.parsing.ingestion_pipeline._EXTRACTORS", {"text/plain": extract}), \
         patch("packages.retrieval.embedding_service.get_embedding_provider", return_value=provider):
        process_document(doc.id, db)
    db.refresh(doc)
    assert doc.status == DocumentStatus.processed
    assert db.query(DocumentChunk).filter_by(document_id=doc.id).one().embedding is not None


def test_embedding_failure_preserves_prior_chunks_and_text(db, test_deal, tmp_path):
    path = tmp_path / "input.txt"
    path.write_text("Updated content", encoding="utf-8")
    doc = document(db, test_deal, path=str(path), text="Previous content")
    db.add(DocumentChunk(document_id=doc.id, chunk_index=0, content="Previous content", page_number=1))
    db.commit()
    provider = SimpleNamespace(embed_texts=lambda _: [])
    with patch("packages.retrieval.embedding_service.get_embedding_provider", return_value=provider):
        with pytest.raises(ValueError, match="Embedding count mismatch"):
            process_document(doc.id, db)
    db.refresh(doc)
    assert doc.status == DocumentStatus.failed
    assert doc.raw_text == "Previous content"
    assert db.query(DocumentChunk).filter_by(document_id=doc.id).one().content == "Previous content"


@pytest.mark.parametrize("status,text", [(DocumentStatus.pending, "Not ready"),
    (DocumentStatus.processing, "Working"), (DocumentStatus.failed, "Old text"),
    (DocumentStatus.processed, "")])
def test_incomplete_corpus_blocks_all_analysis_and_retains_report(db, test_deal, status, text):
    document(db, test_deal, status=status, text=text)
    old_report = Report(deal_id=test_deal.id, title="Previous", content="Previous report",
                        report_type=ReportType.due_diligence, status=ReportStatus.completed)
    db.add(old_report)
    db.commit()

    class NeverCalled(FakeLLMProvider):
        def complete(self, *args, **kwargs):
            raise AssertionError("Incomplete corpus must not call the LLM")

    result = run_deal_analysis_pipeline(test_deal.id, db, provider=NeverCalled())
    assert not result["completed"]
    assert all("error" in stage for stage in result["stages"].values())
    assert "ingestion is incomplete" in result["stages"]["intake"]["error"]
    db.refresh(old_report)
    assert old_report.content == "Previous report"
    assert old_report.status == ReportStatus.completed


def test_batches_cover_every_character_with_overlap_and_bounds():
    text = "".join(f"Riga {i}: clausola àèìòù\n" for i in range(1400))
    batches = list(text_batches(text))
    covered = [False] * len(text)
    for batch in batches:
        assert batch.text == text[batch.start:batch.end]
        assert len(batch.text) <= 8000
        covered[batch.start:batch.end] = [True] * (batch.end - batch.start)
    assert all(covered)
    assert all(right.start < left.end for left, right in zip(batches, batches[1:]))


@pytest.mark.parametrize("kind", [cls.AGENT_NAME for cls in SPECIALISTS] + ["property_identity"])
def test_specialists_analyze_tail_past_old_prefix_limits(kind):
    prompts = []

    class TailOnly(FakeLLMProvider):
        def complete(self, prompt, max_tokens=2048):
            prompts.append(prompt)
            if "TAIL_FINDING_250000" not in prompt:
                return "[]"
            return json.dumps([dict(title="Tail finding", description="TAIL_FINDING_250000",
                finding_type="other", severity="high", confidence=0.9,
                indirizzo="Via Coda 1", foglio="250000")])

    text = "Front matter. " * 2400 + "TAIL_FINDING_250000"
    provider = TailOnly()
    if kind == "property_identity":
        items = PropertyIdentityAgent(provider=provider)._units(1, "visura.txt", text)
    else:
        agent = next(cls for cls in SPECIALISTS if cls.AGENT_NAME == kind)(provider=provider)
        items = collect_document_payloads(text, lambda batch: agent._payloads(1, "doc.txt", batch),
                                          max_chars=agent.BATCH_CHARS)
    assert len(prompts) >= 3
    assert len(items) == 1
    assert items[0]["source_spans"][-1]["end"] == len(text)


def test_overlap_merges_identical_findings_but_keeps_different_amounts():
    def analyze(text):
        return [{"title": "Debt", "amount": 100}, {"title": "Debt", "amount": 200}]
    items = collect_document_payloads("x" * 17000, analyze)
    assert len(items) == 2
    assert {i["amount"] for i in items} == {100, 200}
    assert all(len(item["source_spans"]) == 3 for item in items)


def test_later_specialist_batch_failure_preserves_previous_persisted_findings(db, test_deal):
    doc = document(db, test_deal, text="Matter. " * 2100)
    db.add(AgentFinding(deal_id=test_deal.id, document_id=doc.id, agent_name="encumbrances_review",
        category="mortgage", title="Previous finding", description="Preserve me", severity="high"))
    db.commit()

    class FailLater(FakeLLMProvider):
        calls = 0
        def complete(self, prompt, max_tokens=2048):
            self.calls += 1
            return "[]" if self.calls == 1 else "[{broken"

    with pytest.raises(LLMResponseError):
        EncumbrancesAgent(provider=FailLater()).run(test_deal.id, db, document_ids={doc.id})
    db.rollback()
    assert db.query(AgentFinding).filter_by(agent_name="encumbrances_review").one().title == "Previous finding"


def test_aggregation_batches_include_every_document_and_repeated_identity():
    context = "\n".join(f"=== Document {i}: doc{i}.txt ===\n" + (f"finding_{i}\n" * 1600)
                        for i in range(1, 5))
    batches = aggregation_batches(context)
    assert all(len(batch) <= 8000 for batch in batches)
    for i in range(1, 5):
        relevant = [batch for batch in batches if f"finding_{i}" in batch]
        assert len(relevant) >= 2
        assert all(f"=== Document {i}:" in batch for batch in relevant)
        assert sum(batch.count(f"finding_{i}") for batch in batches) == 1600


def test_aggregator_does_not_hold_connection_during_llm(db, test_deal):
    class CheckConnection(FakeLLMProvider):
        def complete(self, prompt, max_tokens=2048):
            assert not db.in_transaction()
            return "[]"
    result = RedFlagAggregationAgent(provider=CheckConnection()).run(test_deal.id, db)
    assert result["red_flags_count"] == 0


def test_later_aggregation_batch_failure_preserves_prior_flags(db, test_deal, monkeypatch):
    db.add(RedFlag(deal_id=test_deal.id, title="Previous flag", issue="Keep this evidence",
                   severity="high", category="guarantee"))
    db.commit()
    context = "=== Document 1: debt.txt ===\n" + "A documented exposure.\n" * 900
    monkeypatch.setattr(RedFlagAggregationAgent, "_build_context_text", lambda *args: context)

    class FailSecond(FakeLLMProvider):
        calls = 0
        def complete(self, prompt, max_tokens=2048):
            self.calls += 1
            return "[]" if self.calls == 1 else "{broken"

    provider = FailSecond()
    with pytest.raises(LLMResponseError, match="batch 2"):
        RedFlagAggregationAgent(provider=provider).run(test_deal.id, db)
    db.rollback()
    assert provider.calls == 3  # batch 1, batch 2, one format retry of batch 2
    assert db.query(RedFlag).one().title == "Previous flag"


def test_long_identity_document_merges_fields_from_first_and_last_parts():
    class Profiles(FakeLLMProvider):
        def complete(self, prompt, max_tokens=2048):
            item = {"foglio": "12", "particella": "345", "confidence": 0.9}
            if "FIRST_MARKER" in prompt:
                item["categoria"] = "A/2"
            if "LAST_MARKER" in prompt:
                item["rendita"] = "Euro 812,00"
            return json.dumps([item])

    text = ("FIRST_MARKER Foglio 12 particella 345\n" + "Other content.\n" * 2200
            + "LAST_MARKER Foglio 12 particella 345")
    items = PropertyIdentityAgent(provider=Profiles())._units(1, "visura.txt", text)
    rows = [SimpleNamespace(id=index, source_document_id=1, source_kind="visura_catastale", **item)
            for index, item in enumerate(items)]
    units = consolidate_property_profiles(rows)
    assert len(units) == 1
    assert units[0].values["categoria"] == "A/2"
    assert units[0].values["rendita"] == "Euro 812,00"


def test_same_titled_flags_keep_different_documents_amounts_and_years():
    original = {"title": "Tax exposure", "category": "tax", "severity": "high",
                "source_document_id": 1, "amount_eur": 10000, "fiscal_years": [2022]}
    flags = [original, {**original, "source_document_id": 2},
             {**original, "amount_eur": 20000}, {**original, "fiscal_years": [2023]},
             {**original, "title": "TAX EXPOSURE!", "severity": "medium"}]
    deduped, removed = RedFlagAggregationAgent(provider=FakeLLMProvider())._deduplicate(flags)
    assert len(deduped) == 4
    assert removed == 1


@pytest.mark.parametrize("lang", ["en", "it"])
def test_large_report_uses_bounded_summary_and_keeps_every_finding(db, test_deal, lang):
    doc = document(db, test_deal)
    test_deal.report_language = lang
    for index in range(60):
        db.add(RedFlag(deal_id=test_deal.id, title=f"FLAG_UNIQUE_{index:03}",
            issue="Evidence explicitly stated in the source. " * 40,
            recommendation=f"ACTION_UNIQUE_{index:03}", severity="high" if index < 20 else "medium",
            category="contract", source_document_id=doc.id))
    db.commit()

    class NoUnboundedSummary(FakeLLMProvider):
        def complete(self, *args, **kwargs):
            raise AssertionError("Large reports must not send an unbounded summary prompt")

    from tests.report_fixtures import complete_report_inputs
    complete_report_inputs(db, test_deal)
    result = generate_due_diligence_report(test_deal.id, db, provider=NoUnboundedSummary())
    report = db.get(Report, result["report_id"])
    assert report.status == ReportStatus.completed
    assert "60" in report.content and "20" in report.content and "40" in report.content
    assert ("summary is abbreviated" if lang == "en" else "riepilogo è abbreviato") in report.content
    for index in range(60):
        assert f"FLAG_UNIQUE_{index:03}" in report.content
        assert f"ACTION_UNIQUE_{index:03}" in report.content
