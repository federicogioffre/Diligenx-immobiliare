"""Tests for the document ingestion pipeline."""
import io
import os
import tempfile
from pathlib import Path
from unittest.mock import MagicMock, patch

import pytest
from docx import Document as DocxDocument
from pypdf import PdfWriter

from packages.parsing.pdf_parser import extract_pdf, PageText, pdf_page_count
from packages.parsing.docx_parser import extract_docx
from packages.parsing.text_parser import extract_txt
from packages.parsing.language import detect_language
from packages.parsing.chunker import chunk_pages, Chunk
from packages.parsing.ingestion_pipeline import process_document

from app.models.document import Document, DocumentStatus
from app.models.document_chunk import DocumentChunk


# ---------------------------------------------------------------------------
# Helpers: create real files in a temp dir
# ---------------------------------------------------------------------------

@pytest.fixture
def tmp_dir(tmp_path):
    return tmp_path


def _make_txt(tmp_dir: Path, content: str, name: str = "test.txt") -> Path:
    p = tmp_dir / name
    p.write_text(content, encoding="utf-8")
    return p


def _make_pdf(tmp_dir: Path, pages: list[str], name: str = "test.pdf") -> Path:
    """Create a real PDF with text on each page."""
    from pypdf import PdfWriter
    from io import BytesIO
    from reportlab.pdfgen import canvas as rl_canvas
    from reportlab.lib.pagesizes import letter

    writer = PdfWriter()
    for text in pages:
        buf = BytesIO()
        c = rl_canvas.Canvas(buf, pagesize=letter)
        # Write text near top of page
        c.drawString(72, 700, text)
        c.showPage()
        c.save()
        buf.seek(0)
        from pypdf import PdfReader
        reader = PdfReader(buf)
        writer.add_page(reader.pages[0])

    path = tmp_dir / name
    with open(path, "wb") as f:
        writer.write(f)
    return path


def _make_docx(tmp_dir: Path, paragraphs: list[str], name: str = "test.docx") -> Path:
    doc = DocxDocument()
    for p in paragraphs:
        doc.add_paragraph(p)
    path = tmp_dir / name
    doc.save(str(path))
    return path


# ---------------------------------------------------------------------------
# Test: TXT extraction
# ---------------------------------------------------------------------------

class TestExtractTxt:
    def test_basic(self, tmp_dir):
        path = _make_txt(tmp_dir, "Hello world.\nSecond line.")
        pages = extract_txt(path)
        assert len(pages) == 1
        assert pages[0].page_number == 1
        assert "Hello world." in pages[0].text

    def test_empty_file(self, tmp_dir):
        path = _make_txt(tmp_dir, "   ")
        pages = extract_txt(path)
        assert pages == []


# ---------------------------------------------------------------------------
# Test: DOCX extraction
# ---------------------------------------------------------------------------

class TestExtractDocx:
    def test_basic(self, tmp_dir):
        path = _make_docx(tmp_dir, ["First paragraph.", "Second paragraph."])
        pages = extract_docx(path)
        assert len(pages) == 1
        assert "First paragraph." in pages[0].text
        assert "Second paragraph." in pages[0].text
        assert pages[0].page_number == 1

    def test_empty_docx(self, tmp_dir):
        path = _make_docx(tmp_dir, [])
        pages = extract_docx(path)
        assert pages == []


# ---------------------------------------------------------------------------
# Test: PDF extraction (needs reportlab)
# ---------------------------------------------------------------------------

class TestExtractPdf:
    @pytest.fixture(autouse=True)
    def _check_reportlab(self):
        pytest.importorskip("reportlab")

    def test_single_page(self, tmp_dir):
        path = _make_pdf(tmp_dir, ["Hello from page one"])
        pages = extract_pdf(path)
        assert len(pages) == 1
        assert pages[0].page_number == 1
        assert "Hello from page one" in pages[0].text

    def test_multi_page(self, tmp_dir):
        path = _make_pdf(tmp_dir, ["Page 1 text", "Page 2 text"])
        pages = extract_pdf(path)
        assert len(pages) == 2
        assert pages[0].page_number == 1
        assert pages[1].page_number == 2

    def test_page_count(self, tmp_dir):
        path = _make_pdf(tmp_dir, ["A", "B", "C"])
        assert pdf_page_count(path) == 3


# ---------------------------------------------------------------------------
# Test: Language detection
# ---------------------------------------------------------------------------

class TestDetectLanguage:
    def test_english(self):
        text = (
            "The quick brown fox jumps over the lazy dog. "
            "This is a contract between the parties for the provision of services."
        )
        assert detect_language(text) == "en"

    def test_italian(self):
        text = (
            "Il presente contratto è stipulato tra le parti per la fornitura "
            "di servizi. La durata del contratto è di un anno dalla data della firma."
        )
        assert detect_language(text) == "it"

    def test_empty(self):
        assert detect_language("") == "en"

    def test_mixed_defaults_correctly(self):
        # Short ambiguous text — should still return a valid result
        result = detect_language("OK")
        assert result in ("en", "it")


# ---------------------------------------------------------------------------
# Test: Chunker
# ---------------------------------------------------------------------------

class TestChunker:
    def test_single_short_page(self):
        pages = [PageText(page_number=1, text="Short text.")]
        chunks = chunk_pages(pages, chunk_size=1000)
        assert len(chunks) == 1
        assert chunks[0].chunk_index == 0
        assert chunks[0].page_number == 1
        assert chunks[0].text == "Short text."

    def test_long_page_produces_multiple_chunks(self):
        # 2500 chars → should produce ~3 chunks with size=1000 overlap=200
        long_text = "word " * 500  # 2500 chars
        pages = [PageText(page_number=1, text=long_text)]
        chunks = chunk_pages(pages, chunk_size=1000, chunk_overlap=200)
        assert len(chunks) >= 2
        # All from page 1
        assert all(c.page_number == 1 for c in chunks)
        # Indices are sequential
        assert [c.chunk_index for c in chunks] == list(range(len(chunks)))

    def test_multi_page_preserves_page_numbers(self):
        pages = [
            PageText(page_number=1, text="Page one content. " * 10),
            PageText(page_number=2, text="Page two content. " * 10),
        ]
        chunks = chunk_pages(pages, chunk_size=100, chunk_overlap=20)
        page_numbers = {c.page_number for c in chunks}
        assert 1 in page_numbers
        assert 2 in page_numbers

    def test_token_count_set(self):
        pages = [PageText(page_number=1, text="Hello world from the chunker test.")]
        chunks = chunk_pages(pages, chunk_size=5000)
        assert chunks[0].token_count > 0

    def test_empty_pages(self):
        chunks = chunk_pages([])
        assert chunks == []


# ---------------------------------------------------------------------------
# Test: Full pipeline (process_document)
# ---------------------------------------------------------------------------

class TestProcessDocument:
    def test_txt_document(self, db, test_deal, tmp_dir):
        content = (
            "This is a test document for the ingestion pipeline. "
            "It contains enough text to verify that the chunker works. "
            "The quick brown fox jumps over the lazy dog. " * 5
        )
        path = _make_txt(tmp_dir, content)

        doc = Document(
            deal_id=test_deal.id,
            filename="test.txt",
            file_path=str(path),
            mime_type="text/plain",
            file_size=len(content),
            status=DocumentStatus.pending,
        )
        db.add(doc)
        db.commit()
        db.refresh(doc)

        process_document(doc.id, db)
        db.refresh(doc)

        assert doc.status == DocumentStatus.processed
        assert doc.language == "en"
        assert doc.raw_text is not None
        assert len(doc.raw_text) > 0

        chunks = db.query(DocumentChunk).filter(
            DocumentChunk.document_id == doc.id
        ).all()
        assert len(chunks) >= 1
        assert chunks[0].chunk_index == 0
        assert chunks[0].page_number == 1
        assert chunks[0].token_count > 0

    def test_italian_document(self, db, test_deal, tmp_dir):
        content = (
            "Il presente contratto è stipulato tra le parti per la fornitura "
            "di servizi professionali. La durata del contratto è di un anno "
            "dalla data della firma. Le condizioni sono le seguenti. " * 5
        )
        path = _make_txt(tmp_dir, content)

        doc = Document(
            deal_id=test_deal.id,
            filename="contratto.txt",
            file_path=str(path),
            mime_type="text/plain",
            file_size=len(content),
            status=DocumentStatus.pending,
        )
        db.add(doc)
        db.commit()
        db.refresh(doc)

        process_document(doc.id, db)
        db.refresh(doc)

        assert doc.language == "it"
        assert doc.status == DocumentStatus.processed

    def test_docx_document(self, db, test_deal, tmp_dir):
        paragraphs = [
            "This is a DOCX document.",
            "It has multiple paragraphs for testing.",
            "The ingestion pipeline should handle this format correctly.",
        ]
        path = _make_docx(tmp_dir, paragraphs)

        doc = Document(
            deal_id=test_deal.id,
            filename="test.docx",
            file_path=str(path),
            mime_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            file_size=os.path.getsize(path),
            status=DocumentStatus.pending,
        )
        db.add(doc)
        db.commit()
        db.refresh(doc)

        process_document(doc.id, db)
        db.refresh(doc)

        assert doc.status == DocumentStatus.processed
        assert "DOCX document" in doc.raw_text

    def test_pdf_document(self, db, test_deal, tmp_dir):
        pytest.importorskip("reportlab")
        path = _make_pdf(tmp_dir, [
            "This is page one of the contract.",
            "This is page two with additional terms.",
        ])

        doc = Document(
            deal_id=test_deal.id,
            filename="contract.pdf",
            file_path=str(path),
            mime_type="application/pdf",
            file_size=os.path.getsize(path),
            status=DocumentStatus.pending,
        )
        db.add(doc)
        db.commit()
        db.refresh(doc)

        process_document(doc.id, db)
        db.refresh(doc)

        assert doc.status == DocumentStatus.processed
        assert doc.page_count == 2
        assert doc.language == "en"

        chunks = db.query(DocumentChunk).filter(
            DocumentChunk.document_id == doc.id
        ).all()
        assert len(chunks) >= 1

    def test_nonexistent_document_raises(self, db):
        with pytest.raises(ValueError, match="not found"):
            process_document(99999, db)

    def test_bad_file_marks_failed(self, db, test_deal, tmp_dir):
        # Point to a nonexistent file
        doc = Document(
            deal_id=test_deal.id,
            filename="missing.pdf",
            file_path=str(tmp_dir / "does_not_exist.pdf"),
            mime_type="application/pdf",
            file_size=0,
            status=DocumentStatus.pending,
        )
        db.add(doc)
        db.commit()
        db.refresh(doc)

        with pytest.raises(Exception):
            process_document(doc.id, db)

        db.refresh(doc)
        assert doc.status == DocumentStatus.failed

    def test_idempotent_rerun(self, db, test_deal, tmp_dir):
        """Running process_document twice should replace old chunks."""
        content = "This is the original content for idempotency test. " * 10
        path = _make_txt(tmp_dir, content)

        doc = Document(
            deal_id=test_deal.id,
            filename="idem.txt",
            file_path=str(path),
            mime_type="text/plain",
            file_size=len(content),
            status=DocumentStatus.pending,
        )
        db.add(doc)
        db.commit()
        db.refresh(doc)

        process_document(doc.id, db)
        first_count = db.query(DocumentChunk).filter(
            DocumentChunk.document_id == doc.id
        ).count()

        # Re-process
        doc.status = DocumentStatus.pending
        db.commit()
        process_document(doc.id, db)
        second_count = db.query(DocumentChunk).filter(
            DocumentChunk.document_id == doc.id
        ).count()

        assert first_count == second_count


# ---------------------------------------------------------------------------
# Test: Upload route triggers ingestion
# ---------------------------------------------------------------------------

class TestUploadIngestion:
    def test_upload_txt_triggers_processing(self, client, auth_headers, test_deal):
        content = b"This is a plain text document for the upload integration test."
        resp = client.post(
            f"/documents/upload?deal_id={test_deal.id}",
            headers=auth_headers,
            files={"file": ("test.txt", io.BytesIO(content), "text/plain")},
        )
        assert resp.status_code == 201
        doc_id = resp.json()["id"]

        # Upload is 201 + background ingestion; verify final state via the
        # list endpoint once the BackgroundTask has run.
        listing = client.get(
            f"/documents/deals/{test_deal.id}/documents", headers=auth_headers
        ).json()
        final = next(d for d in listing if d["id"] == doc_id)
        assert final["status"] == "processed"
        assert final["language"] == "en"

    def test_upload_docx_triggers_processing(self, client, auth_headers, test_deal, tmp_dir):
        docx_doc = DocxDocument()
        docx_doc.add_paragraph("Il presente documento è un test di integrazione.")
        buf = io.BytesIO()
        docx_doc.save(buf)
        buf.seek(0)

        resp = client.post(
            f"/documents/upload?deal_id={test_deal.id}",
            headers=auth_headers,
            files={
                "file": (
                    "test.docx",
                    buf,
                    "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                )
            },
        )
        assert resp.status_code == 201
        doc_id = resp.json()["id"]

        listing = client.get(
            f"/documents/deals/{test_deal.id}/documents", headers=auth_headers
        ).json()
        final = next(d for d in listing if d["id"] == doc_id)
        assert final["status"] == "processed"
        assert final["language"] == "it"


# ---------------------------------------------------------------------------
# Test: OCR fallback for scanned PDFs
# ---------------------------------------------------------------------------

class TestOCRFallback:
    """Tests for OCR fallback on scanned/image-only PDFs."""

    @staticmethod
    def _make_blank_pdf(tmp_dir: Path, n_pages: int = 1) -> Path:
        """Create a PDF with blank pages (no embedded text)."""
        writer = PdfWriter()
        for _ in range(n_pages):
            writer.add_blank_page(width=612, height=792)
        path = tmp_dir / "scanned.pdf"
        with open(path, "wb") as f:
            writer.write(f)
        return path

    def test_ocr_called_for_blank_page(self, tmp_dir):
        """extract_pdf calls _ocr_page when a page has no embedded text."""
        path = self._make_blank_pdf(tmp_dir)
        with patch("packages.parsing.pdf_parser._ocr_page", return_value="Testo OCR.") as mock_ocr:
            pages = extract_pdf(path)
        mock_ocr.assert_called_once_with(str(path), 1)
        assert len(pages) == 1
        assert pages[0].text == "Testo OCR."
        assert pages[0].page_number == 1

    def test_ocr_not_called_for_text_page(self, tmp_dir):
        """extract_pdf does NOT call _ocr_page when the page has embedded text."""
        pytest.importorskip("reportlab")
        path = _make_pdf(tmp_dir, ["Embedded text on page one."])
        with patch("packages.parsing.pdf_parser._ocr_page") as mock_ocr:
            pages = extract_pdf(path)
        mock_ocr.assert_not_called()
        assert len(pages) == 1

    def test_ocr_empty_result_fails_document(self, tmp_dir):
        """Unreadable pages must never be silently excluded from analysis."""
        path = self._make_blank_pdf(tmp_dir)
        with patch("packages.parsing.pdf_parser._ocr_page", return_value="   "):
            with pytest.raises(ValueError, match="PDF page 1.*incomplete"):
                extract_pdf(path)

    def test_ocr_mixed_pdf(self, tmp_dir):
        """Mixed PDF: native text pages use pypdf, blank pages use OCR."""
        pytest.importorskip("reportlab")
        text_path = _make_pdf(tmp_dir, ["Embedded text page one"])

        from pypdf import PdfWriter as _PW, PdfReader as _PR
        writer = _PW()
        with open(text_path, "rb") as f:
            reader = _PR(f)
            for page in reader.pages:
                writer.add_page(page)
        writer.add_blank_page(width=612, height=792)
        merged = tmp_dir / "mixed.pdf"
        with open(merged, "wb") as f:
            writer.write(f)

        called_for = []

        def side_effect(pdf_path, page_number):
            called_for.append(page_number)
            return "OCR page two"

        with patch("packages.parsing.pdf_parser._ocr_page", side_effect=side_effect):
            pages = extract_pdf(merged)

        assert called_for == [2]   # OCR only for the blank page
        assert len(pages) == 2
        assert "Embedded text page one" in pages[0].text
        assert pages[1].text == "OCR page two"

    def test_ocr_unavailable_returns_empty(self, tmp_dir):
        """_ocr_page returns '' when OCR libraries are not installed."""
        path = self._make_blank_pdf(tmp_dir)
        from packages.parsing.pdf_parser import _ocr_page
        with patch("packages.parsing.pdf_parser._OCR_AVAILABLE", False):
            result = _ocr_page(str(path), 1)
        assert result == ""

    def test_ocr_page_exception_returns_empty(self, tmp_dir):
        """_ocr_page returns '' when pdf2image/tesseract raises an exception."""
        path = self._make_blank_pdf(tmp_dir)
        from packages.parsing.pdf_parser import _ocr_page
        with patch("packages.parsing.pdf_parser._OCR_AVAILABLE", True):
            with patch(
                "packages.parsing.pdf_parser._pdf2image_convert",
                side_effect=Exception("Poppler not found"),
            ):
                result = _ocr_page(str(path), 1)
        assert result == ""

    def test_ocr_page_uses_ita_eng_language(self, tmp_dir):
        """_ocr_page calls pytesseract with lang='ita+eng'."""
        path = self._make_blank_pdf(tmp_dir)
        mock_image = MagicMock()
        mock_tess = MagicMock()
        mock_tess.image_to_string.return_value = "testo"

        from packages.parsing.pdf_parser import _ocr_page
        with patch("packages.parsing.pdf_parser._OCR_AVAILABLE", True):
            with patch("packages.parsing.pdf_parser._pdf2image_convert", return_value=[mock_image]):
                with patch("packages.parsing.pdf_parser._pytesseract", mock_tess):
                    result = _ocr_page(str(path), 1)

        mock_tess.image_to_string.assert_called_once_with(mock_image, lang="ita+eng")
        assert result == "testo"

    def test_ocr_page_uses_dpi_300(self, tmp_dir):
        """_ocr_page rasterizes at 300 DPI."""
        path = self._make_blank_pdf(tmp_dir)
        mock_convert = MagicMock(return_value=[MagicMock()])
        mock_tess = MagicMock()
        mock_tess.image_to_string.return_value = "text"

        from packages.parsing.pdf_parser import _ocr_page
        with patch("packages.parsing.pdf_parser._OCR_AVAILABLE", True):
            with patch("packages.parsing.pdf_parser._pdf2image_convert", mock_convert):
                with patch("packages.parsing.pdf_parser._pytesseract", mock_tess):
                    _ocr_page(str(path), 1)

        mock_convert.assert_called_once_with(str(path), first_page=1, last_page=1, dpi=300)

    def test_ocr_page_number_passed_correctly(self, tmp_dir):
        """_ocr_page passes the correct page number to convert_from_path."""
        path = self._make_blank_pdf(tmp_dir, n_pages=3)
        mock_convert = MagicMock(return_value=[MagicMock()])
        mock_tess = MagicMock()
        mock_tess.image_to_string.return_value = "text"

        from packages.parsing.pdf_parser import _ocr_page
        with patch("packages.parsing.pdf_parser._OCR_AVAILABLE", True):
            with patch("packages.parsing.pdf_parser._pdf2image_convert", mock_convert):
                with patch("packages.parsing.pdf_parser._pytesseract", mock_tess):
                    _ocr_page(str(path), 3)

        mock_convert.assert_called_once_with(str(path), first_page=3, last_page=3, dpi=300)
